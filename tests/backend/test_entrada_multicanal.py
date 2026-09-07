from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document

from src.backend.entrada_multicanal import (
    DocumentoOrigem,
    DocxExtractionError,
    EntradaMulticanalError,
    criar_documento_extensao,
    criar_documento_texto,
    extrair_documento_docx,
    extrair_documento_pdf,
)


def test_documento_origem_aceita_quatro_canais():
    pdf = DocumentoOrigem(canal="pdf", texto="texto do pdf", identificador_origem="arquivo.pdf", tamanho_bytes=10)
    docx = DocumentoOrigem(canal="docx", texto="texto do docx", identificador_origem="arquivo.docx", tamanho_bytes=20)
    texto = DocumentoOrigem(canal="texto", texto="texto colado", identificador_origem="clipboard", tamanho_bytes=15)
    extensao = DocumentoOrigem(canal="extensao", texto="texto da extensão", identificador_origem="campo-1", tamanho_bytes=25)

    assert {pdf.canal, docx.canal, texto.canal, extensao.canal} == {"pdf", "docx", "texto", "extensao"}


def test_entrada_multicanal_excecoes_herdam_da_base():
    assert issubclass(DocxExtractionError, EntradaMulticanalError)


def test_extrair_documento_pdf_wraps_pdf_existing_text():
    with patch("src.backend.entrada_multicanal.extract_pdf_text", return_value="texto extraido"):
        documento = extrair_documento_pdf("arquivo.pdf")

    assert documento.canal == "pdf"
    assert documento.texto == "texto extraido"
    assert documento.identificador_origem == "arquivo.pdf"


def test_extrair_documento_docx_com_paragrafos(tmp_path):
    arquivo = tmp_path / "documento.docx"
    document = Document()
    document.add_paragraph("Primeiro parágrafo")
    document.add_paragraph("Segundo parágrafo")
    document.save(arquivo)

    documento = extrair_documento_docx(arquivo)

    assert documento.canal == "docx"
    assert "Primeiro parágrafo" in documento.texto
    assert "Segundo parágrafo" in documento.texto


def test_extrair_documento_docx_rejeita_arquivo_invalido(tmp_path):
    arquivo = tmp_path / "invalido.docx"
    arquivo.write_bytes(b"nao-e-um-docx")

    with pytest.raises(DocxExtractionError):
        extrair_documento_docx(arquivo)


def test_criar_documento_texto_rejeita_texto_acima_do_limite():
    with pytest.raises(ValueError):
        criar_documento_texto("x" * (10 * 1024 * 1024 + 1), max_size_bytes=10 * 1024 * 1024)


def test_criar_documento_extensao_preserva_identificador():
    documento = criar_documento_extensao("texto da extensao", "campo-123")

    assert documento.canal == "extensao"
    assert documento.identificador_origem == "campo-123"
    assert documento.texto == "texto da extensao"
